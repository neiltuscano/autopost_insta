"""
f1jobs — Excel Job Tracker Builder
Creates and updates f1jobs_tracker.xlsx with all fetched jobs.
Columns: Date, Company, Title, Location, Salary Min/Max, Salary Display,
         Type, H1B, Apply URL, Source, Status, Card Generated, Posted Date
"""
from __future__ import annotations  # Python 3.8 compat

import os
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.filters import AutoFilter
from datetime import datetime

from config import TRACKER_PATH

# ── Brand colors (openpyxl uses ARGB hex) ────────────────────────────────────
BG_DARK    = "FF08080B"   # near-black bg (not used in Excel theme but for ref)
ORANGE     = "FFFF5F00"   # f1jobs orange
ORANGE_LIGHT="FFFF8C42"
WHITE      = "FFFFFFFF"
DARK_GRAY  = "FF1A1A2E"
MID_GRAY   = "FF2E2E4E"
LIGHT_GRAY = "FFB0B7C3"
GREEN_BG   = "FFD6F5E3"
GREEN_FG   = "FF00A651"
YELLOW_BG  = "FFFFF9C4"
YELLOW_FG  = "FFF57F17"
RED_BG     = "FFFFE0E0"
RED_FG     = "FFD32F2F"
HEADER_BG  = "FF0D0F17"
ROW_ALT    = "FFF5F7FF"
ROW_NORM   = "FFFFFFFF"

COLUMNS = [
    ("Date Fetched",    16),
    ("Company",         22),
    ("Job Title",       34),
    ("City",            16),
    ("State",           8),
    ("Salary Min ($)",  16),
    ("Salary Max ($)",  16),
    ("Salary Display",  18),
    ("Job Type",        14),
    ("H-1B Sponsored",  16),
    ("Apply URL",       40),
    ("Source",          12),
    ("Status",          14),
    ("Card Made",       12),
    ("Posted Date",     14),
    ("Notes",           30),
]


def _thin_border():
    s = Side(style="thin", color="FFD0D5E0")
    return Border(left=s, right=s, top=s, bottom=s)


def _header_font():
    return Font(name="Arial", bold=True, color=WHITE, size=10)


def _data_font(bold=False, color="FF1A1A2E"):
    return Font(name="Arial", bold=bold, color=color, size=10)


def _center():
    return Alignment(horizontal="center", vertical="center", wrap_text=False)


def _left():
    return Alignment(horizontal="left", vertical="center", wrap_text=False)


def create_or_open_tracker() -> openpyxl.Workbook:
    """Load existing tracker or create a fresh one."""
    if os.path.exists(TRACKER_PATH):
        wb = openpyxl.load_workbook(TRACKER_PATH)
        print(f"  Opened existing tracker: {TRACKER_PATH}")
        return wb

    wb = openpyxl.Workbook()

    # ── Jobs sheet ────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Jobs"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A3"

    # Row 1 — Title banner
    ws.merge_cells("A1:P1")
    title_cell = ws["A1"]
    title_cell.value = "f1jobs · F-1 & H-1B Job Tracker · USA"
    title_cell.font  = Font(name="Arial", bold=True, color=WHITE, size=14)
    title_cell.fill  = PatternFill("solid", fgColor=ORANGE)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    # Row 2 — Column headers
    ws.row_dimensions[2].height = 24
    for col_idx, (header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.font      = _header_font()
        cell.fill      = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = _center()
        cell.border    = _thin_border()
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # ── Dashboard sheet ───────────────────────────────────────────────────────
    ds = wb.create_sheet("Dashboard")
    ds.sheet_view.showGridLines = False

    ds.merge_cells("A1:F1")
    ds["A1"].value     = "f1jobs · Daily Dashboard"
    ds["A1"].font      = Font(name="Arial", bold=True, color=WHITE, size=16)
    ds["A1"].fill      = PatternFill("solid", fgColor=ORANGE)
    ds["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ds.row_dimensions[1].height = 36

    stats = [
        ("Total Jobs Tracked",  '=COUNTA(Jobs!B3:B10000)'),
        ("Jobs Posted",         '=COUNTIF(Jobs!M3:M10000,"Posted")'),
        ("Jobs Pending",        '=COUNTIF(Jobs!M3:M10000,"Pending")'),
        ("Cards Generated",     '=COUNTIF(Jobs!N3:N10000,"Yes")'),
        ("H-1B Confirmed",      '=COUNTIF(Jobs!J3:J10000,"Yes")'),
        ("Unique Companies",    '=SUMPRODUCT(1/COUNTIF(Jobs!B3:B10000,Jobs!B3:B10000))'),
    ]
    for row, (label, formula) in enumerate(stats, start=3):
        ds.cell(row=row, column=1, value=label).font  = _data_font(bold=True)
        ds.cell(row=row, column=1).fill               = PatternFill("solid", fgColor=ROW_ALT)
        ds.cell(row=row, column=1).alignment          = _left()
        val_cell = ds.cell(row=row, column=2, value=formula)
        val_cell.font      = _data_font(bold=True, color=ORANGE)
        val_cell.alignment = _center()
        ds.column_dimensions["A"].width = 24
        ds.column_dimensions["B"].width = 16

    wb.active = ws
    print(f"  Created new tracker: {TRACKER_PATH}")
    return wb


def _existing_keys(ws) -> set:
    """Collect company+title keys already in the sheet (rows 3+)."""
    keys = set()
    for row in ws.iter_rows(min_row=3, max_col=3, values_only=True):
        company, title = (row[1] or ""), (row[2] or "")
        if company or title:
            keys.add(f"{company.lower().strip()}|{title.lower().strip()}")
    return keys


def _location_parts(location: str):
    """Split 'Dallas, TX' → ('Dallas', 'TX')."""
    parts = [p.strip() for p in location.split(",")]
    if len(parts) >= 2:
        return parts[0], parts[1]
    return parts[0], ""


def _status_style(cell, status):
    if status == "Posted":
        cell.fill = PatternFill("solid", fgColor=GREEN_BG)
        cell.font = _data_font(bold=True, color=GREEN_FG)
    elif status == "Pending":
        cell.fill = PatternFill("solid", fgColor=YELLOW_BG)
        cell.font = _data_font(color=YELLOW_FG)
    elif status == "Skipped":
        cell.fill = PatternFill("solid", fgColor=RED_BG)
        cell.font = _data_font(color=RED_FG)


def append_jobs(jobs: list[dict], wb: openpyxl.Workbook) -> int:
    """
    Append new jobs to the Jobs sheet. Skips duplicates.
    Returns count of new rows added.
    """
    ws = wb["Jobs"]
    existing = _existing_keys(ws)

    # Find next empty row
    next_row = 3
    for row in ws.iter_rows(min_row=3, max_col=2):
        if any(cell.value for cell in row):
            next_row += 1
        else:
            break

    added = 0
    for job in jobs:
        key = f"{job.get('company','').lower().strip()}|{job.get('title','').lower().strip()}"
        if key in existing:
            continue
        existing.add(key)

        city, state = _location_parts(job.get("location", ""))
        is_alt      = (next_row % 2 == 0)
        row_fill    = PatternFill("solid", fgColor=ROW_ALT if is_alt else ROW_NORM)

        values = [
            job.get("date_fetched", datetime.now().strftime("%Y-%m-%d")),
            job.get("company", ""),
            job.get("title", ""),
            city,
            state,
            job.get("salary_min", None),
            job.get("salary_max", None),
            job.get("salary", "Competitive"),
            job.get("type", "Full-time"),
            "Yes" if job.get("h1b") else "No",
            job.get("apply_url", ""),
            job.get("source", ""),
            "Pending",   # Status
            "No",        # Card Made
            "",          # Posted Date
            "",          # Notes
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            cell.fill      = row_fill
            cell.border    = _thin_border()
            cell.alignment = _left()
            cell.font      = _data_font()
            ws.row_dimensions[next_row].height = 20

            # Special styling
            if col_idx == 2:  # Company — bold orange
                cell.font = _data_font(bold=True, color=ORANGE)
            elif col_idx == 3:  # Title — bold
                cell.font = _data_font(bold=True)
            elif col_idx in (6, 7):  # Salary numbers — right align
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0'
            elif col_idx == 8:  # Salary display — green bold
                cell.font = _data_font(bold=True, color=GREEN_FG)
            elif col_idx == 10:  # H1B — center green
                cell.alignment = _center()
                if val == "Yes":
                    cell.font = _data_font(bold=True, color=GREEN_FG)
            elif col_idx == 13:  # Status
                _status_style(cell, val)
                cell.alignment = _center()
            elif col_idx == 14:  # Card Made
                cell.alignment = _center()

        next_row += 1
        added   += 1

    # Auto-filter on header row
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{next_row - 1}"
    return added


def mark_cards_generated(jobs: list[dict], wb: openpyxl.Workbook):
    """Mark 'Card Made' = Yes for jobs that had cards generated today."""
    ws  = wb["Jobs"]
    today = datetime.now().strftime("%Y-%m-%d")
    targets = {
        f"{j.get('company','').lower().strip()}|{j.get('title','').lower().strip()}"
        for j in jobs
    }
    for row in ws.iter_rows(min_row=3):
        company = str(row[1].value or "").lower().strip()
        title   = str(row[2].value or "").lower().strip()
        if f"{company}|{title}" in targets:
            row[13].value = "Yes"   # Card Made (col 14, index 13)
            row[13].font  = _data_font(bold=True, color=GREEN_FG)
            row[13].alignment = Alignment(horizontal="center", vertical="center")


def save_tracker(wb: openpyxl.Workbook):
    wb.save(TRACKER_PATH)
    print(f"  Tracker saved → {TRACKER_PATH}")


def build_tracker(jobs: list[dict]) -> str:
    """Full flow: open/create tracker, append jobs, save. Returns path."""
    wb    = create_or_open_tracker()
    added = append_jobs(jobs, wb)
    print(f"  Added {added} new jobs to tracker")
    save_tracker(wb)
    return TRACKER_PATH


if __name__ == "__main__":
    from fetcher import get_daily_jobs
    print("\nBuilding tracker...\n")
    jobs = get_daily_jobs(10)
    path = build_tracker(jobs)
    print(f"\n  ✓ Tracker ready: {path}")
