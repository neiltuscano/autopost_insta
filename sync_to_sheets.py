"""
f1jobs — Google Sheets Sync
Pushes all jobs from f1jobs_tracker.xlsx to a public Google Sheet.

First run: creates the sheet automatically and saves the URL to .env
Subsequent runs: finds the sheet by ID from .env and overwrites the data.

Usage:
  python sync_to_sheets.py           # sync xlsx → Google Sheets
  python sync_to_sheets.py --dry-run # show what would be pushed
"""
from __future__ import annotations

import os, re, sys, argparse
from pathlib import Path
from datetime import datetime

import openpyxl
import gspread
from google.oauth2.service_account import Credentials

# ── Config ────────────────────────────────────────────────────────────────────

def _load_env():
    env_path = Path(__file__).parent / ".env"
    if env_path.exists():
        for line in env_path.read_text().splitlines():
            line = line.strip()
            if line and not line.startswith("#") and "=" in line:
                k, _, v = line.partition("=")
                os.environ.setdefault(k.strip(), v.strip())

_load_env()

ROOT          = Path(__file__).parent
KEY_PATH      = ROOT / "f1jobs-sheets-key.json"
TRACKER_PATH  = ROOT / "data" / "f1jobs_tracker.xlsx"
SHEET_TITLE   = "f1jobs · F-1 & H-1B Job Tracker · USA"
SCOPES        = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]

COLUMNS = [
    "Date Fetched", "Company", "Job Title", "City", "State",
    "Salary Min ($)", "Salary Max ($)", "Salary Display", "Job Type",
    "H-1B Sponsored", "Apply URL", "Source", "Status", "Card Made",
    "Posted Date", "Notes",
]

# ── Helpers ───────────────────────────────────────────────────────────────────

def _gc() -> gspread.Client:
    if not KEY_PATH.exists():
        raise FileNotFoundError(
            f"Service account key not found: {KEY_PATH}\n"
            "  Download it from Google Cloud Console → IAM → Service Accounts → Keys"
        )
    creds = Credentials.from_service_account_file(str(KEY_PATH), scopes=SCOPES)
    return gspread.authorize(creds)


def _update_env(key: str, value: str):
    """Insert or replace a key=value line in .env."""
    env_path = ROOT / ".env"
    content  = env_path.read_text() if env_path.exists() else ""
    pattern  = rf"^{re.escape(key)}=.*$"
    new_line = f"{key}={value}"
    if re.search(pattern, content, re.MULTILINE):
        content = re.sub(pattern, new_line, content, flags=re.MULTILINE)
    else:
        content = content.rstrip("\n") + "\n" + new_line + "\n"
    env_path.write_text(content)


def _get_or_create_sheet(gc: gspread.Client) -> gspread.Spreadsheet:
    """Find the sheet by saved ID, or create a new one."""
    sheet_id = os.getenv("GOOGLE_SHEET_ID", "")

    if sheet_id:
        try:
            sh = gc.open_by_key(sheet_id)
            print(f"  Found existing sheet: {sh.url}")
            return sh
        except gspread.SpreadsheetNotFound:
            print(f"  Sheet ID {sheet_id[:8]}... not found — creating new one.")

    # Create new sheet
    sh = gc.create(SHEET_TITLE)
    print(f"  Created new sheet: {sh.url}")

    # Make publicly readable (anyone with link can view)
    sh.client.request(
        "POST",
        f"https://www.googleapis.com/drive/v3/files/{sh.id}/permissions",
        json={"role": "reader", "type": "anyone"},
    )
    print("  Set to public (anyone with link can view)")

    # Save ID + URL to .env for future runs
    _update_env("GOOGLE_SHEET_ID",  sh.id)
    _update_env("GOOGLE_SHEET_URL", sh.url)
    print(f"  Saved GOOGLE_SHEET_ID and GOOGLE_SHEET_URL to .env")

    return sh


def _read_xlsx() -> list[list]:
    """Read all job rows from the xlsx tracker. Returns list of rows."""
    if not TRACKER_PATH.exists():
        raise FileNotFoundError(f"Tracker not found: {TRACKER_PATH}")

    wb = openpyxl.load_workbook(TRACKER_PATH, data_only=True)
    ws = wb["Jobs"]

    rows = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        # Skip completely empty rows
        if not any(v for v in row):
            continue
        # Convert all values to strings (Sheets API wants strings)
        clean = []
        for v in row:
            if v is None:
                clean.append("")
            elif isinstance(v, float) and v == int(v):
                clean.append(str(int(v)))
            else:
                clean.append(str(v))
        rows.append(clean)

    return rows


# ── Main sync ─────────────────────────────────────────────────────────────────

def sync(dry_run: bool = False) -> str:
    """
    Push xlsx data to Google Sheets. Returns the sheet URL.
    """
    print(f"\n{'='*60}")
    print(f"  f1jobs Sheets Sync — {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    print(f"{'='*60}")

    # Read local tracker
    print("\n  [1/3]  Reading xlsx tracker…")
    rows = _read_xlsx()
    print(f"  Found {len(rows)} job rows")

    if dry_run:
        print("\n  [DRY RUN] Would push these columns:")
        print("  " + " | ".join(COLUMNS))
        print(f"\n  [DRY RUN] Would push {len(rows)} rows")
        if rows:
            print(f"  [DRY RUN] First row: {rows[0][:4]}")
        print("\n  [DRY RUN] No changes made.")
        return ""

    # Connect + get/create sheet
    print("\n  [2/3]  Connecting to Google Sheets…")
    gc = _gc()
    sh = _get_or_create_sheet(gc)
    ws = sh.sheet1

    # Clear and rewrite
    print("\n  [3/3]  Uploading data…")
    ws.clear()

    # Header row
    ws.update("A1", [COLUMNS])

    # Format header row
    ws.format("A1:P1", {
        "backgroundColor": {"red": 0.05, "green": 0.06, "blue": 0.09},
        "textFormat": {"bold": True, "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
        "horizontalAlignment": "CENTER",
    })

    # Data rows
    if rows:
        ws.update(f"A2:P{len(rows)+1}", rows)

    # Freeze header row
    sh.sheet1.freeze(rows=1)

    url = sh.url
    print(f"\n  ✓  Synced {len(rows)} jobs to Google Sheets")
    print(f"  URL: {url}")
    print()

    return url


# ── CLI ───────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="f1jobs → Google Sheets sync")
    parser.add_argument("--dry-run", action="store_true",
                        help="Preview without making any changes")
    args = parser.parse_args()
    sync(dry_run=args.dry_run)


if __name__ == "__main__":
    main()
